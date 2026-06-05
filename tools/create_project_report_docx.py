from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "reports"
OUT_PATH = OUT_DIR / "JR Shelf Life Label App - Detailed Project Report.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "DADCE0", size: str = "4") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa: int = 9360, indent_dxa: int = 0) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def style_run(run, *, bold=False, italic=False, size=None, color="000000") -> None:
    run.bold = bold
    run.italic = italic
    if size:
        run.font.size = Pt(size)
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.color.rgb = None


def add_paragraph(doc: Document, text: str = "", style: str | None = None):
    paragraph = doc.add_paragraph(style=style)
    if text:
        run = paragraph.add_run(text)
        style_run(run)
    return paragraph


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    run = paragraph.add_run(text)
    style_run(run)


def add_numbered(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Number")
    run = paragraph.add_run(text)
    style_run(run)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_width(table)

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_width(cell, widths[idx])
        set_cell_shading(cell, "F8F9FA")
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)
        run = cell.paragraphs[0].add_run(header)
        style_run(run, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_width(cell, widths[idx])
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.paragraphs[0].paragraph_format.space_after = Pt(0)
            run = cell.paragraphs[0].add_run(value)
            style_run(run)

    doc.add_paragraph()


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    title.font.size = Pt(24)
    title.font.bold = True
    title.paragraph_format.space_after = Pt(10)

    for name, size, before, after in (
        ("Heading 1", 16, 18, 8),
        ("Heading 2", 14, 14, 6),
        ("Heading 3", 12, 10, 4),
    ):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)


def add_title_page(doc: Document) -> None:
    title = doc.add_paragraph(style="Title")
    title.add_run("JR Shelf Life Label App").bold = True
    subtitle = doc.add_paragraph()
    subtitle_run = subtitle.add_run("Detailed Project Report for Presentation and Handover")
    style_run(subtitle_run, size=14)
    metadata = doc.add_paragraph()
    metadata_run = metadata.add_run("Prepared for: JR Foods project stakeholders | Date: 28 May 2026 | Status: Android Phase 1 functional build")
    style_run(metadata_run, size=10)
    doc.add_paragraph()
    add_paragraph(
        doc,
        "This report summarizes the business purpose, current implementation, technical architecture, data model, printing workflow, current APK status, risks, and recommended next steps for the JR Shelf Life Label App.",
    )
    doc.add_page_break()


def build_report() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    doc.add_heading("1. Executive Summary", level=1)
    add_paragraph(
        doc,
        "The JR Shelf Life Label App is an Android-first mobile application built for central kitchen and outlet teams to print accurate secondary shelf-life labels. It standardizes product shelf-life rules, outlet-specific product visibility, label preview, quantity selection, Bluetooth printer selection, and local print history in a simple offline workflow.",
    )
    add_paragraph(
        doc,
        "The project has progressed from requirements and planning into a functional Android app with a real APK installed and tested on an Oppo phone. The app can select outlets, browse product categories, manage local product data, preview labels, print via a GPrinter/Gainscha Bluetooth bridge, and export/import local data through Excel.",
    )
    add_table(
        doc,
        ["Area", "Current Status", "Presentation Notes"],
        [
            ["Business workflow", "Functional Phase 1 staff flow", "Outlet selection, product/category browsing, preview, quantity, and print are implemented."],
            ["Android app", "Working APK available", "Latest release APK is built and installed successfully on the test phone."],
            ["Printing", "Functional Android Bluetooth printing", "Labels are rendered as bitmaps and sent through a native GPrinter SDK bridge."],
            ["Admin tools", "Broad local admin coverage", "Admin can manage products, categories, outlets, outlet lists, printers, and import/export."],
            ["Production readiness", "Prototype-to-pilot stage", "Main remaining hardening items are SQLite persistence, validation, and release workflow."],
        ],
        [1800, 2300, 5260],
    )

    doc.add_heading("2. Business Problem and Objective", level=1)
    add_paragraph(
        doc,
        "Food teams need to print accurate shelf-life labels quickly and consistently. Manual handwritten labels or manual date calculations create avoidable operational risk: wrong expiry timing, inconsistent label format, unclear storage method, and slower service during busy periods.",
    )
    add_paragraph(
        doc,
        "The app solves this by giving staff a fixed operational flow: choose outlet, choose product, confirm label preview, choose quantity, and print. Shelf-life timing is calculated automatically from product rules, and the printed label uses a standardized format.",
    )
    doc.add_heading("Target Users", level=2)
    for item in [
        "Central kitchen staff who prepare or open food items.",
        "Outlet staff who need fast product lookup and printing.",
        "Outlet managers who control local product visibility.",
        "Local company admin users who maintain product/category/outlet records.",
        "Future site and company admins who will manage cloud data in Phase 2.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("3. Scope and Phase Position", level=1)
    add_paragraph(
        doc,
        "Phase 1 is focused on an offline Android app. It deliberately avoids cloud login, multi-company SaaS administration, and full web portal work so the team can validate the staff workflow and printer behavior first.",
    )
    add_table(
        doc,
        ["Scope Category", "Included in Current Phase", "Deferred"],
        [
            ["Staff printing workflow", "Outlet, category, product detail, label preview, quantity, Bluetooth printing", "Cloud user roles and cross-device sync"],
            ["Local admin", "Categories, products, outlets, outlet product lists, printer choice, import/export", "Server-backed master data approval flow"],
            ["Data", "Seed data, local persistence, Excel import/export, bundled images", "PostgreSQL/cloud sync, audit-grade history"],
            ["Platforms", "Android APK", "iOS Bluetooth printing and App Store/TestFlight release"],
        ],
        [2100, 3900, 3360],
    )

    doc.add_heading("4. Current Application Flow", level=1)
    for step in [
        "User opens the app and selects an outlet.",
        "App shows available categories for that outlet.",
        "User chooses a category and sees the product list.",
        "User opens a product detail page.",
        "App shows product information, shelf-life rule, label preview, and quantity controls.",
        "User chooses Normal or Compact label preview mode.",
        "User taps Print; app captures the label preview as a bitmap and sends it to the selected Bluetooth printer.",
        "App saves a local PrintJob record for history and export.",
    ]:
        add_numbered(doc, step)

    doc.add_heading("5. Feature Breakdown", level=1)
    add_table(
        doc,
        ["Feature", "Implemented", "Details"],
        [
            ["Outlet selection", "Yes", "User must choose outlet before product printing."],
            ["Category selection", "Yes", "Pastel color-coded category cards with product counts."],
            ["Product browsing", "Yes", "Compact and regular product list modes; search support."],
            ["Product detail", "Yes", "Product image, English/Chinese names, storage method, shelf-life, preview, quantity, print."],
            ["Normal/Compact label preference", "Yes", "User choice persists across products, printing, and app restart."],
            ["Timestamp printing", "Yes", "Separate timestamp label flow for 40 x 30 labels."],
            ["Admin PIN", "Yes", "Simple local access control for admin screens."],
            ["Admin products/categories/outlets", "Yes", "Local create/edit workflows are implemented."],
            ["Outlet product lists", "Yes", "Controls which products are visible per outlet."],
            ["Printer selection", "Yes", "Admin can select preferred Bluetooth label printer."],
            ["Import/export", "Yes", "Excel workbook import/export for operational data."],
            ["Print history", "Basic", "Shows what was printed, when, and quantity."],
        ],
        [2400, 1400, 5560],
    )

    doc.add_heading("6. UI and UX Direction", level=1)
    add_paragraph(
        doc,
        "The UI follows an Operational Clean direction. The design is intentionally plain, readable, and fast for kitchen/outlet staff. The app avoids decorative screens and keeps the main path focused on finding products and printing labels.",
    )
    for item in [
        "Portrait-only layout for consistency on staff phones.",
        "Category-first navigation to reduce filter clutter on the product page.",
        "Compact product mode for seeing more products at once.",
        "Regular product mode for image-based recognition.",
        "Sticky back headers on scroll-heavy admin and detail screens.",
        "Admin access hidden behind long press/PIN behavior to avoid accidental data edits.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("7. Data Model", level=1)
    add_paragraph(
        doc,
        "The current TypeScript domain model defines the core records needed for Phase 1. The planned documentation recommends SQLite, while the current app still uses seed data and SecureStore-style local persistence for many records. That is acceptable for prototype and pilot testing, but SQLite should be introduced before production rollout.",
    )
    add_table(
        doc,
        ["Entity", "Purpose", "Key Fields"],
        [
            ["Outlet", "Represents central kitchen or outlet location", "id, name, code, shortCode"],
            ["ProductCategory", "Groups products and controls product code prefix", "id, name, codePrefix, color, scope"],
            ["Product", "Printable food item with shelf-life rule", "productCode, nameEn, nameZh, daysToExpire, hoursToExpire, storageMethod, imageUri, active"],
            ["OutletProduct", "Outlet-specific product visibility", "outletId, productId, active"],
            ["PrintJob", "Local print history record", "labelType, outletName, productCode, quantity, printedAt, expiresAt, printerName"],
            ["PreferredBluetoothPrinter", "Saved printer preference", "address, name"],
        ],
        [1900, 3300, 4160],
    )

    doc.add_heading("8. Label and Printing Architecture", level=1)
    add_paragraph(
        doc,
        "The app renders labels as on-screen React Native views, captures them as PNG bitmaps, and sends the bitmap to the printer through an Android native module. This avoids relying on the printer's built-in Chinese fonts and gives the app more control over visual layout.",
    )
    add_table(
        doc,
        ["Label Type", "Size", "Current Use"],
        [
            ["Normal product label", "60 x 40 mm", "Full product label with product name, Chinese name/storage, open time, use-by time."],
            ["Compact product label", "40 x 30 mm", "Condensed label with OPEN and USE BY lines; mode preference is saved."],
            ["Timestamp label", "40 x 30 mm", "Standalone date/time label for simple marking workflows."],
        ],
        [2300, 1800, 5260],
    )
    add_paragraph(
        doc,
        "The Android native module uses the GPrinter SDK's BluetoothPort and LabelCommand. It opens a fresh Bluetooth connection, sends a high-contrast bitmap, waits briefly, then closes the port to reduce stuck-connection problems.",
    )

    doc.add_heading("9. Printer Management", level=1)
    for item in [
        "Admin > Printers shows only supported label printers, not general Bluetooth accessories.",
        "The app filters devices by known printer names and Android Bluetooth imaging class.",
        "Admin can choose a preferred printer when multiple printers are paired.",
        "Printing uses the saved preferred printer first.",
        "If the preferred printer is unavailable, the app falls back to the first supported printer.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("10. Import, Export, and APK Variants", level=1)
    add_paragraph(
        doc,
        "The app supports Excel-based import/export to help admins manage data during Phase 1. Export includes categories, outlets, products, outlet product settings, print jobs, and export metadata.",
    )
    add_table(
        doc,
        ["Build Artifact", "Purpose", "Status"],
        [
            ["jr-shelf-life-label-app-android-release-2026-05-28.apk", "Latest Android product-data build", "Built and installed successfully on phone."],
            ["jr-shelf-life-label-app-android-empty-products-2026-05-26.apk", "APK with categories but no seeded products", "Built for clean-data distribution/testing."],
        ],
        [3700, 3400, 2260],
    )

    doc.add_heading("11. Current Technical Risks", level=1)
    add_table(
        doc,
        ["Risk", "Impact", "Recommended Action"],
        [
            ["SecureStore used for broad data persistence", "May not scale well for larger product/history data", "Move structured data to SQLite."],
            ["Printer device filtering is heuristic", "Some real printer models may be hidden if names/classes differ", "Add observed model names after field testing."],
            ["Failed print attempts not fully represented", "History may not capture all operational failures", "Save failed PrintJob records with error message."],
            ["Docs contain garbled Chinese encoding", "Presentation and handover clarity issue", "Normalize docs to UTF-8 and repair Chinese text."],
            ["iOS printing limitations", "Apple rollout may require different Bluetooth strategy", "Research CoreBluetooth/MFi requirements after Android pilot."],
            ["Large uncommitted worktree", "Higher risk of losing track of changes", "Commit current stable version to GitHub."],
        ],
        [2700, 3200, 3460],
    )

    doc.add_heading("12. Roadmap", level=1)
    add_table(
        doc,
        ["Priority", "Next Step", "Why It Matters"],
        [
            ["1", "Commit the stable Android build to GitHub", "Creates a reliable restore point before deeper refactors."],
            ["2", "Move local data to SQLite", "Improves reliability for product records, history, and import/export."],
            ["3", "Harden validation rules", "Prevents duplicate product codes and invalid shelf-life values."],
            ["4", "Improve print failure tracking", "Supports troubleshooting and operational accountability."],
            ["5", "Run multi-printer field test", "Validates preferred printer selection and filtering."],
            ["6", "Prepare pilot release package", "Makes it easier for outlet staff to test the app consistently."],
            ["7", "Plan cloud Phase 2", "Defines server admin, sync, users, and audit logs."],
        ],
        [1200, 3600, 4560],
    )

    doc.add_heading("13. Conclusion", level=1)
    add_paragraph(
        doc,
        "The JR Shelf Life Label App is now a functional Android Phase 1 implementation rather than only a concept. It covers the core staff workflow, local admin operations, label preview, Bluetooth printing, and APK distribution. The project is ready for controlled operational testing, with the main engineering focus now shifting from feature creation to persistence hardening, validation, print reliability, and release management.",
    )

    doc.save(OUT_PATH)


if __name__ == "__main__":
    build_report()
    print(OUT_PATH)
